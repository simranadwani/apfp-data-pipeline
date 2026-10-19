#!/usr/bin/env python3
"""
Reference converter: reads ONLY APFP_Lineage_Spec.xlsx (sheets Pages, Shapes, Shape_Edges) and writes a draw.io file.
It shows how the sheets map to draw.io XML; any other tool can follow the same steps.
Usage: python3 rebuild_from_xlsx.py APFP_Lineage_Spec.xlsx out.drawio
"""
import sys
from xml.sax.saxutils import escape, quoteattr
from openpyxl import load_workbook


def rows(ws):
    it = ws.iter_rows(values_only=True)
    head = [h for h in next(it)]
    return [dict(zip(head, r)) for r in it if any(v is not None for v in r)]


def main(src, dst):
    wb = load_workbook(src, data_only=True)
    pages, shapes, edges = rows(wb['Pages']), rows(wb['Shapes']), rows(wb['Shape_Edges'])
    out = ['<mxfile host="app.diagrams.net" version="24.0.0">']
    for pg in pages:
        cells = ['<mxCell id="0"/>', '<mxCell id="1" parent="0"/>']
        for s in sorted([s for s in shapes if s['page'] == pg['page']], key=lambda s: s['z']):
            val = quoteattr(escape(s['text'] or '').replace('\n', '<br>'))
            st = quoteattr(s['drawio_style'])
            parent = quoteattr(s['parent_id'] or '1')
            if s['type'] == 'TableRow':
                geom = f'<mxGeometry y="{s["rel_y"]}" width="{s["width"]}" height="{s["height"]}" as="geometry"/>'
            elif s['type'] == 'TableCell':
                xa = f' x="{s["rel_x"]}"' if s['rel_x'] else ''
                geom = (f'<mxGeometry{xa} width="{s["width"]}" height="{s["height"]}" as="geometry"><mxRectangle width="{s["width"]}" '
                        f'height="{s["height"]}" as="alternateBounds"/></mxGeometry>')
            else:
                geom = f'<mxGeometry x="{s["x"]}" y="{s["y"]}" width="{s["width"]}" height="{s["height"]}" as="geometry"/>'
            if s['type'] in ('Table', 'TableRow'):
                props = f' layer={quoteattr(s["layer"] or "")} table={quoteattr(s["table"] or "")}' + (f' column={quoteattr(s["column"])}' if s['column'] else '')
                lab = quoteattr(s['label_html'] or '') if s['type'] == 'Table' else '""'
                cells.append(f'<object label={lab} id={quoteattr(s["shape_id"])}{props}><mxCell style={st} vertex="1" parent={parent}>{geom}</mxCell></object>')
            else:
                cells.append(f'<mxCell id={quoteattr(s["shape_id"])} value={val} style={st} vertex="1" parent={parent}>{geom}</mxCell>')
        for e in [e for e in edges if e['page'] == pg['page']]:
            pts = f'<Array as="points"><mxPoint x="{e["lane_x"]}" y="{e["source_y"]}"/></Array>' if e['waypoint'] == 'yes' else ''
            cells.append(f'<mxCell id={quoteattr(e["edge_id"])} style={quoteattr(e["drawio_style"])} edge="1" parent="1" source={quoteattr(e["source_id"])} '
                         f'target={quoteattr(e["target_id"])}><mxGeometry relative="1" as="geometry">{pts}</mxGeometry></mxCell>')
        model = (f'<mxGraphModel dx="1400" dy="900" grid="1" gridSize="10" guides="1" tooltips="1" connect="1" arrows="1" fold="1" page="0" '
                 f'pageScale="1" pageWidth="{pg["width"]}" pageHeight="{pg["height"]}" math="0" shadow="0"><root>{"".join(cells)}</root></mxGraphModel>')
        out.append(f'<diagram id="{pg["diagram_id"]}" name={quoteattr(pg["page"])}>{model}</diagram>')
    out.append('</mxfile>')
    open(dst, 'w', encoding='utf-8').write(''.join(out))
    print('wrote', dst)


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
